from openpyxl import load_workbook
from typing import Tuple, List, Dict
from fastapi import HTTPException


def read_excel(path: str) -> Tuple[object, List[str]]:
    """Carga el workbook y retorna el objeto workbook y la lista de hojas."""
    try:
        wb = load_workbook(filename=path, data_only=True)
        return wb, wb.sheetnames
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error leyendo Excel: {str(e)}")


def read_sheet(wb, sheet_name: str) -> Tuple[List[str], List[Dict]]:
    """
    Retorna (headers_normalizados, data_as_list_of_dicts).

    Normaliza encabezados a minúsculas y sin espacios (espacios -> _).
    """
    if sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail="Hoja no encontrada en el workbook")

    sheet = wb[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))

    if not rows:
        return [], []

    # Encabezados crudos
    raw_headers = [str(h) if h is not None else "" for h in rows[0]]
    headers = [h.strip().lower().replace(" ", "_") for h in raw_headers]

    data = []
    for row in rows[1:]:
        # Saltar filas completamente vacías
        if all(cell is None for cell in row):
            continue

        row_dict = {
            headers[i]: row[i] if i < len(row) else None
            for i in range(len(headers))
        }
        data.append(row_dict)

    return headers, data
